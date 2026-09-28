import importlib.util
import json
import shlex
import subprocess

from bunbird.agent_skills import SOURCE, instructions, parser_manifest


def test_managed_skill_snapshot_mount_and_disable(tmp_path):
    import base64

    from bunbird.agent_skills import managed_instructions

    skill = {
        "id": "12345678-1234-1234-1234-123456789abc",
        "revision": "a" * 64,
        "name": "Example",
        "files": {
            "SKILL.md": base64.b64encode(b"# Example").decode(),
            "scripts/helper.py": base64.b64encode(b"print('hello')").decode(),
        },
    }
    text = managed_instructions(tmp_path, [skill])
    mounted = tmp_path / "managed-skills" / skill["id"] / skill["revision"]
    assert str((mounted / "SKILL.md").as_posix()) in text
    assert (mounted / "scripts/helper.py").read_text() == "print('hello')"
    (mounted / "SKILL.md").write_text("modified by previous task")
    managed_instructions(tmp_path, [skill])
    assert (mounted / "SKILL.md").read_text() == "# Example"
    disabled = managed_instructions(tmp_path, [])
    assert "Example" not in disabled and "不要加载" in disabled


def test_managed_skill_rejects_traversal(tmp_path):
    import base64

    import pytest

    from bunbird.agent_skills import managed_instructions

    for name in ("../escape", "C:/escape", "/escape", "scripts/../../escape", "scripts\\escape"):
        skill = {
            "id": "12345678-1234-1234-1234-123456789abc",
            "revision": "a" * 64,
            "name": "Example",
            "files": {"SKILL.md": "Iw==", name: base64.b64encode(b"bad").decode()},
        }
        with pytest.raises(ValueError, match="Invalid skill path"):
            managed_instructions(tmp_path, [skill])


spec = importlib.util.spec_from_file_location("attachment_parser", SOURCE / "scripts/parse_file.py")
parser = importlib.util.module_from_spec(spec)
spec.loader.exec_module(parser)


def test_mounted_skill_runs_and_preserves_source(tmp_path):
    folder = tmp_path / "session with spaces"
    folder.mkdir()
    source = folder / "数据 'test.csv"
    original = "编号,值\n001,=1+2\n002,中文\n".encode("gb18030")
    source.write_bytes(original)
    assert instructions(folder) == ""
    manifest = parser_manifest(folder, source)
    subprocess.run(shlex.split(manifest["parseCommand"]), check=True, capture_output=True)
    result = json.loads((source.parent / (source.name + ".parsed.json")).read_text("utf-8"))
    assert result["blocks"][1]["content"] == ["001", "=1+2"]
    assert source.read_bytes() == original
    assert "read-attachments" in instructions(folder)
    assert parser_manifest(folder, folder / "a.exe") == {"parser": "unsupported"}


def test_excel_formats_and_pagination(tmp_path):
    import openpyxl
    import xlwt

    path = tmp_path / "test.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["001", "=1+2"])
    book.active.append(["second", 4])
    book.create_sheet("其他").append(["中文"])
    book.save(path)
    result = parser.extract(path, limit=1)
    assert result["status"] == "partial"
    assert len(result["sheets"]) == 2
    assert result["blocks"][0]["content"][1] == {"cell": "B1", "value": "=1+2", "cached": None}
    assert parser.extract(path, offset=1)["blocks"][0]["content"][0]["value"] == "second"
    path = tmp_path / "old.xls"
    book = xlwt.Workbook()
    book.add_sheet("数据").write(0, 0, "旧格式")
    book.save(str(path))
    assert parser.extract(path)["blocks"][0]["content"][0] == {"cell": "A1", "value": "旧格式"}


def test_word_order_and_pdf_text_and_password(tmp_path):
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    path = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("前文")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "表格"
    doc.add_paragraph("后文")
    doc.save(path)
    assert [b["content"] for b in parser.extract(path)["blocks"]] == ["前文", [["表格"]], "后文"]
    path = tmp_path / "test.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 100 Td (Hello PDF) Tj ET")
    page[NameObject("/Contents")] = stream
    writer.add_blank_page(width=200, height=200)
    writer.write(path)
    result = parser.extract(path)
    assert "Hello PDF" in result["blocks"][0]["content"]
    assert any("OCR" in w for w in result["warnings"])
    writer.encrypt("password")
    writer.write(path)
    assert parser.extract(path)["error"] == "PASSWORD_REQUIRED"


def test_bad_input_and_budget_are_explicit(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not pdf")
    assert parser.extract(path)["error"] == "PARSE_FAILED"
    path = tmp_path / "large.txt"
    path.write_text("a" * 100001, encoding="utf-8")
    assert parser.extract(path)["status"] == "partial"
    outcome = subprocess.run(
        [
            __import__("sys").executable,
            str(SOURCE / "scripts/parse_file.py"),
            "--input",
            str(path),
            "--output",
            str(path),
        ],
        capture_output=True,
    )
    assert outcome.returncode != 0
    assert path.stat().st_size == 100001
