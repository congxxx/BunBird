"""Bounded, read-only extraction. Standalone so the skill can run inside a session."""

import argparse
import csv
import io
import json
import logging
import zipfile
from pathlib import Path

logging.getLogger("pypdf").setLevel(logging.CRITICAL)


class Limit(Exception):
    pass


class Extract:
    def __init__(self):
        self.result = {"status": "ok", "warnings": [], "truncated": False, "blocks": []}
        self.remaining = 100_000

    def add(self, location, content):
        text = json.dumps(content, ensure_ascii=False, default=str)
        if len(text) > self.remaining:
            self.result["truncated"] = True
            raise Limit()
        self.remaining -= len(text)
        self.result["blocks"].append({"location": location, "content": content})

    def warn(self, message):
        if message not in self.result["warnings"]:
            self.result["warnings"].append(message)


def check_archive(path):
    # Inspect declared sizes before OOXML libraries decompress the archive.
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024:
            raise ValueError("ARCHIVE_TOO_LARGE")


def decode(path):
    data = path.read_bytes()
    encodings = (
        ["utf-16"] if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "gb18030"]
    )
    for encoding in encodings:
        try:
            return data.decode(encoding), encoding
        except UnicodeError:
            pass
    raise ValueError("TEXT_ENCODING_UNKNOWN")


def extract(path, offset=0, limit=200):
    out = Extract()
    result = out.result
    result.update(format=path.suffix.lower(), offset=offset, limit=limit)
    try:
        if path.stat().st_size > 10 * 1024 * 1024:
            raise ValueError("FILE_TOO_LARGE")
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            import openpyxl

            check_archive(path)
            book = openpyxl.load_workbook(path, read_only=True, keep_links=False)
            cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
            result["sheets"] = []
            out.warn("公式未重新计算；缓存值可能缺失或过期。图表和图片未提取。")
            try:
                for sheet in book:
                    result["sheets"].append(
                        {
                            "name": sheet.title,
                            "state": sheet.sheet_state,
                            "rows": sheet.max_row,
                            "columns": sheet.max_column,
                        }
                    )
                    result["truncated"] |= sheet.max_row > offset + limit or sheet.max_column > 100
                    rows = sheet.iter_rows(
                        min_row=offset + 1,
                        max_row=min(sheet.max_row, offset + limit),
                        max_col=min(100, sheet.max_column),
                    )
                    values = cached[sheet.title].iter_rows(
                        min_row=offset + 1,
                        max_row=min(sheet.max_row, offset + limit),
                        max_col=min(100, sheet.max_column),
                    )
                    if offset >= sheet.max_row:
                        continue
                    for row, cache_row in zip(rows, values, strict=True):
                        cells = [
                            {
                                "cell": cell.coordinate,
                                "value": cell.value,
                                **({"cached": cache.value} if cell.data_type == "f" else {}),
                            }
                            for cell, cache in zip(row, cache_row, strict=True)
                            if cell.value is not None
                        ]
                        if cells:
                            out.add(sheet.title, cells)
            finally:
                book.close()
                cached.close()
        elif suffix == ".xls":
            import xlrd

            book = xlrd.open_workbook(path, on_demand=True)
            result["sheets"] = []
            out.warn("读取 xls 保存的值，不还原公式、图表或图片。")
            try:
                for sheet in book.sheets():
                    result["sheets"].append(
                        {"name": sheet.name, "rows": sheet.nrows, "columns": sheet.ncols}
                    )
                    result["truncated"] |= sheet.nrows > offset + limit or sheet.ncols > 100
                    for row in range(offset, min(sheet.nrows, offset + limit)):
                        cells = []
                        for col in range(min(sheet.ncols, 100)):
                            cell = sheet.cell(row, col)
                            value = (
                                xlrd.xldate.xldate_as_datetime(
                                    cell.value, book.datemode
                                ).isoformat()
                                if cell.ctype == xlrd.XL_CELL_DATE
                                else cell.value
                            )
                            cells.append({"cell": xlrd.formula.cellname(row, col), "value": value})
                        out.add(sheet.name, cells)
            finally:
                book.release_resources()
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ValueError("PASSWORD_REQUIRED")
            result["pages"] = len(reader.pages)
            result["truncated"] = len(reader.pages) > offset + limit
            out.warn("仅提取文字，不保证表格关系或视觉版式准确。")
            for i in range(offset, min(len(reader.pages), offset + limit)):
                text = reader.pages[i].extract_text() or ""
                if not text.strip():
                    out.warn(f"第 {i + 1} 页未提取到文字，需要 OCR 或视觉核对。")
                out.add(f"page {i + 1}", text)
        elif suffix == ".docx":
            from docx import Document
            from docx.table import Table

            check_archive(path)
            document = Document(path)
            out.warn("只提取正文段落及表格；不含页眉页脚、文本框、图片、批注及修订。")
            for i, block in enumerate(document.iter_inner_content()):
                if i < offset:
                    continue
                if i >= offset + limit:
                    result["truncated"] = True
                    break
                if isinstance(block, Table):
                    result["truncated"] |= len(block.rows) > 500 or len(block.columns) > 100
                    out.add(
                        f"block {i + 1} table",
                        [[cell.text for cell in row.cells[:100]] for row in block.rows[:500]],
                    )
                else:
                    out.add(f"block {i + 1} paragraph", block.text)
        elif suffix in {".csv", ".tsv", ".txt", ".md", ".json"}:
            text, encoding = decode(path)
            result["encoding"] = encoding
            if suffix in {".csv", ".tsv"}:
                try:
                    delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
                except csv.Error:
                    delimiter = "\t" if suffix == ".tsv" else ","
                result["delimiter"] = delimiter
                rows = csv.reader(io.StringIO(text), delimiter=delimiter)
            else:
                rows = iter(text.splitlines())
            for i, row in enumerate(rows):
                if i < offset:
                    continue
                if i >= offset + limit:
                    result["truncated"] = True
                    break
                out.add(f"row {i + 1}", row)
        else:
            raise ValueError("UNSUPPORTED_FORMAT")
    except Limit:
        out.warn("达到正文字符预算，结果已截断；请缩小读取区间。")
    except Exception as error:
        allowed = {
            "FILE_TOO_LARGE",
            "ARCHIVE_TOO_LARGE",
            "PASSWORD_REQUIRED",
            "TEXT_ENCODING_UNKNOWN",
            "UNSUPPORTED_FORMAT",
        }
        result["status"] = "error"
        result["error"] = str(error) if str(error) in allowed else "PARSE_FAILED"
        out.warn("未完成解析，不能把已提取片段当作完整文件。")
    if result["truncated"] and result["status"] == "ok":
        result["status"] = "partial"
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--limit", type=int, default=200)
    args = parser.parse_args()
    if (
        args.offset < 0
        or not 1 <= args.limit <= 500
        or args.input.resolve() == args.output.resolve()
    ):
        parser.error("invalid range or output would overwrite input")
    result = extract(args.input, args.offset, args.limit)
    args.output.write_text(
        json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    print(
        json.dumps(
            {
                "status": result["status"],
                "blocks": len(result["blocks"]),
                "output": str(args.output),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
